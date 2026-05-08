from __future__ import annotations

import datetime as dt
import glob
import json
import math
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from time import perf_counter
from types import SimpleNamespace
from typing import Any, Optional, Sequence

import numpy as np
import pandas as pd

from shared_lib import format_duration
from shared_lib.cli_token_lists import parse_comma_separated_cli_tokens
from shared_lib.results import EXIT_ERROR, EXIT_VALIDATION_ERROR, SharedResult, failure, success

try:
    import fitz  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    fitz = None

try:
    import pdfplumber  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    pdfplumber = None

try:
    import camelot  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    camelot = None

try:
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    Image = None

try:
    import pytesseract  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    pytesseract = None

try:
    import easyocr  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    easyocr = None

try:
    from docx import Document  # type: ignore[import-not-found]
except Exception:  # noqa: BLE001
    Document = None


SUPPORTED_INPUT = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_FORMATS = {"txt", "json", "docx"}
TABLE_FORMATS = {"csv", "xlsx", "json", "docx"}

OCR_PDF_QUALITY_PRESETS: dict[str, dict[str, Any]] = {
    "premium": {"oversample": 450, "optimize": 3, "jpg_quality": 95, "png_quality": 95, "pagesegmode": 1},
    "balanced": {"oversample": 350, "optimize": 2, "jpg_quality": 85, "png_quality": 85, "pagesegmode": 3},
    "fast": {"oversample": 250, "optimize": 1, "jpg_quality": 80, "png_quality": 80, "pagesegmode": 6},
}


@dataclass
class ExtractedPage:
    source_file: str
    page: int
    text: str


@dataclass
class ExtractedTable:
    source_file: str
    page: int
    table_index: int
    dataframe: pd.DataFrame


def _stamp() -> str:
    return dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")


def _timestamped_path(input_path: Path, mode: str, ext: str, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / f"{input_path.stem}_{mode}_{_stamp()}.{ext}"


def _resolve_input_paths(raw_inputs: Sequence[str], *, allow_empty_match: bool, max_files: int) -> list[Path]:
    paths: list[Path] = []
    for raw in raw_inputs:
        text = str(raw or "").strip()
        if not text:
            continue
        expanded = str(Path(text).expanduser())
        has_glob = glob.has_magic(expanded)
        matches = [Path(m).resolve() for m in glob.glob(expanded, recursive=True)] if has_glob else []
        if has_glob:
            if not matches and not allow_empty_match:
                raise FileNotFoundError(f"No matches for pattern: {text}")
            paths.extend(matches)
            continue
        p = Path(expanded).resolve()
        if not p.exists():
            raise FileNotFoundError(f"Input path not found: {p}")
        paths.append(p)

    out: list[Path] = []
    seen: set[Path] = set()
    for p in paths:
        if p.is_file() and p.suffix.lower() in SUPPORTED_INPUT:
            r = p.resolve()
            if r not in seen:
                seen.add(r)
                out.append(r)
    if len(out) > max_files:
        raise ValueError(f"Resolved input file count {len(out)} exceeds max_files={max_files}.")
    return out


def _extract_pdf_text_layout(pdf_path: Path) -> list[ExtractedPage]:
    if fitz is None:
        raise RuntimeError("Missing dependency: pymupdf")
    doc = fitz.open(pdf_path)
    pages: list[ExtractedPage] = []
    try:
        for idx in range(len(doc)):
            page = doc.load_page(idx)
            text_dict = page.get_text("dict")
            blocks = text_dict.get("blocks", [])
            blocks = sorted(
                blocks,
                key=lambda b: (
                    round(float(b.get("bbox", [0, 0, 0, 0])[1]), 1),
                    round(float(b.get("bbox", [0, 0, 0, 0])[0]), 1),
                ),
            )
            lines: list[str] = []
            for b in blocks:
                if b.get("type", 0) != 0:
                    continue
                for line in b.get("lines", []):
                    spans = line.get("spans", [])
                    text = "".join(str(s.get("text", "")) for s in spans).strip()
                    if text:
                        lines.append(text)
                if lines and lines[-1] != "":
                    lines.append("")
            pages.append(ExtractedPage(source_file=str(pdf_path), page=idx + 1, text="\n".join(lines).strip()))
    finally:
        doc.close()
    return pages


def _text_chars_count(pages: Sequence[ExtractedPage]) -> int:
    return int(sum(len(str(p.text or "").replace("\n", "").strip()) for p in pages))


def _ocr_pdf(input_pdf: Path, out_pdf: Path, *, lang: str, quality: str, output_type: str) -> tuple[bool, str]:
    if shutil.which("ocrmypdf") is None:
        return False, "ocrmypdf executable not found in PATH"
    preset = OCR_PDF_QUALITY_PRESETS.get(quality, OCR_PDF_QUALITY_PRESETS["premium"])
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    cmd = ["ocrmypdf"]
    if shutil.which("unpaper") is not None:
        cmd.extend(["--clean", "--clean-final"])
    optimize_level = int(preset["optimize"])
    # ocrmypdf optimize levels 2/3 may require pngquant; fall back so OCR can still run.
    if optimize_level >= 2 and shutil.which("pngquant") is None:
        optimize_level = 1
    cmd.extend(
        [
            "--deskew",
            "--rotate-pages",
            "--tesseract-oem",
            "1",
            "--tesseract-pagesegmode",
            str(preset["pagesegmode"]),
            "--oversample",
            str(preset["oversample"]),
            "--optimize",
            str(optimize_level),
            "--jpg-quality",
            str(preset["jpg_quality"]),
            "--png-quality",
            str(preset["png_quality"]),
            "--output-type",
            output_type,
            "-l",
            lang,
            "--force-ocr",
            str(input_pdf),
            str(out_pdf),
        ]
    )
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True)
        return True, ""
    except subprocess.CalledProcessError as exc:
        stderr = str(exc.stderr or "").strip()
        stdout = str(exc.stdout or "").strip()
        detail = stderr or stdout or f"ocrmypdf exited with code {exc.returncode}"
        return False, detail
    except Exception as exc:  # noqa: BLE001
        return False, str(exc)


def _image_to_temp_pdf(image_path: Path, *, dpi: int) -> Path:
    if Image is None:
        raise RuntimeError("Missing dependency: pillow")
    temp_pdf = Path(tempfile.gettempdir()) / f"{image_path.stem}_img_{_stamp()}.pdf"
    with Image.open(image_path) as im:
        rgb = im.convert("RGB")
        rgb.save(temp_pdf, "PDF", resolution=float(dpi))
    return temp_pdf


def _maybe_ocr_pdf(
    pdf_path: Path,
    ocr_mode: str,
    threshold: int,
    *,
    lang: str,
    quality: str,
    output_type: str,
) -> Path:
    if ocr_mode == "never":
        return pdf_path
    if ocr_mode == "always":
        ocr_out = Path(tempfile.gettempdir()) / f"{pdf_path.stem}_ocr_{_stamp()}.pdf"
        ok, _ = _ocr_pdf(pdf_path, ocr_out, lang=lang, quality=quality, output_type=output_type)
        return ocr_out if ok else pdf_path

    pages = _extract_pdf_text_layout(pdf_path)
    chars = _text_chars_count(pages)
    if chars < threshold:
        ocr_out = Path(tempfile.gettempdir()) / f"{pdf_path.stem}_ocr_{_stamp()}.pdf"
        ok, _ = _ocr_pdf(pdf_path, ocr_out, lang=lang, quality=quality, output_type=output_type)
        return ocr_out if ok else pdf_path
    return pdf_path


def _extract_tables_camelot(pdf_path: Path) -> list[ExtractedTable]:
    if camelot is None:
        return []
    out: list[ExtractedTable] = []
    page_counter: dict[int, int] = {}
    for flavor in ("lattice", "stream"):
        try:
            found = camelot.read_pdf(str(pdf_path), pages="all", flavor=flavor)
        except Exception:
            continue
        for table in found:
            df = table.df
            if df is None or df.empty or len(df) < 2:
                continue
            try:
                page = int(str(table.page))
            except Exception:
                page = 1
            page_counter[page] = page_counter.get(page, 0) + 1
            out.append(
                ExtractedTable(
                    source_file=str(pdf_path),
                    page=page,
                    table_index=page_counter[page],
                    dataframe=df.reset_index(drop=True),
                )
            )
    return out


def _extract_tables_pdfplumber(pdf_path: Path) -> list[ExtractedTable]:
    if pdfplumber is None:
        return []
    out: list[ExtractedTable] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            raw_tables = page.extract_tables() or []
            table_idx = 0
            for raw in raw_tables:
                if not raw or len(raw) < 2:
                    continue
                header = [str(x) if x is not None else "" for x in raw[0]]
                rows = raw[1:]
                if not rows:
                    continue
                df = pd.DataFrame(rows, columns=header)
                if df.empty:
                    continue
                table_idx += 1
                out.append(ExtractedTable(source_file=str(pdf_path), page=page_idx, table_index=table_idx, dataframe=df))
    return out


def _extract_pdf_tables(pdf_path: Path, engine: str) -> list[ExtractedTable]:
    if engine == "camelot":
        return _extract_tables_camelot(pdf_path)
    if engine == "pdfplumber":
        return _extract_tables_pdfplumber(pdf_path)
    tables = _extract_tables_camelot(pdf_path)
    if tables:
        return tables
    return _extract_tables_pdfplumber(pdf_path)


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except Exception:
        return default


def _otsu_threshold(gray: np.ndarray) -> int:
    hist, _ = np.histogram(gray.ravel(), bins=256, range=(0, 256))
    total = gray.size
    sum_total = float(np.dot(np.arange(256), hist))
    sum_bg = 0.0
    weight_bg = 0.0
    max_between = -1.0
    threshold = 128
    for t in range(256):
        weight_bg += float(hist[t])
        if weight_bg <= 0:
            continue
        weight_fg = float(total) - weight_bg
        if weight_fg <= 0:
            break
        sum_bg += float(t * hist[t])
        mean_bg = sum_bg / weight_bg
        mean_fg = (sum_total - sum_bg) / weight_fg
        between = weight_bg * weight_fg * (mean_bg - mean_fg) ** 2
        if between > max_between:
            max_between = between
            threshold = t
    return int(threshold)


def _estimate_skew_angle(binary_img: Image.Image) -> float:
    arr = np.array(binary_img)
    ys, xs = np.where(arr < 128)
    if xs.size < 200:
        return 0.0
    x = xs.astype(np.float64)
    y = ys.astype(np.float64)
    x -= x.mean()
    y -= y.mean()
    cov = np.cov(np.vstack((x, y)))
    evals, evecs = np.linalg.eig(cov)
    principal = evecs[:, int(np.argmax(evals))]
    angle = math.degrees(math.atan2(float(principal[1]), float(principal[0])))
    if angle > 90:
        angle -= 180
    if angle < -90:
        angle += 180
    # ignore unrealistic values
    if abs(angle) > 20:
        return 0.0
    return float(angle)


def _crop_text_zone_tesseract(image: Image.Image, *, lang: str) -> Image.Image:
    if pytesseract is None:
        return image
    try:
        data = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT, config="--psm 6")
    except Exception:
        return image

    n = len(data.get("text", []))
    boxes: list[tuple[int, int, int, int]] = []
    for i in range(n):
        text = str(data["text"][i] or "").strip()
        conf = _safe_int(data.get("conf", [0] * n)[i], default=0)
        if not text or conf < 30:
            continue
        x = _safe_int(data.get("left", [0] * n)[i], default=0)
        y = _safe_int(data.get("top", [0] * n)[i], default=0)
        w = _safe_int(data.get("width", [0] * n)[i], default=0)
        h = _safe_int(data.get("height", [0] * n)[i], default=0)
        if w > 2 and h > 2:
            boxes.append((x, y, x + w, y + h))

    if not boxes:
        return image

    x0 = max(0, min(b[0] for b in boxes) - 20)
    y0 = max(0, min(b[1] for b in boxes) - 20)
    x1 = min(image.width, max(b[2] for b in boxes) + 20)
    y1 = min(image.height, max(b[3] for b in boxes) + 20)
    if x1 - x0 < 50 or y1 - y0 < 50:
        return image
    return image.crop((x0, y0, x1, y1))


def _preprocess_for_ocr(image: Image.Image, *, lang: str) -> Image.Image:
    # 1) grayscale + denoise
    gray = ImageOps.grayscale(image)
    gray = gray.filter(ImageFilter.MedianFilter(size=3))
    # 2) local contrast + slight sharpening
    gray = ImageOps.autocontrast(gray, cutoff=1)
    gray = ImageEnhance.Contrast(gray).enhance(1.8)
    gray = ImageEnhance.Sharpness(gray).enhance(1.3)

    # 3) binarize with Otsu threshold
    arr = np.array(gray, dtype=np.uint8)
    thr = _otsu_threshold(arr)
    bin_arr = np.where(arr > thr, 255, 0).astype(np.uint8)
    binary = Image.fromarray(bin_arr, mode="L")

    # 4) deskew (small-angle)
    angle = _estimate_skew_angle(binary)
    if abs(angle) > 0.2:
        binary = binary.rotate(-angle, expand=True, fillcolor=255)

    # 5) crop to text-heavy region using tesseract boxes
    cropped = _crop_text_zone_tesseract(binary, lang=lang)
    return cropped


def _score_ocr_text(text: str) -> int:
    return sum(2 for ch in text if ch.isalnum()) + sum(1 for ch in text if ch.strip())


def _ocr_image_tesseract(image_path: Path, lang: str) -> str:
    if Image is None or pytesseract is None:
        raise RuntimeError("Missing dependencies for tesseract OCR: pillow + pytesseract")
    with Image.open(image_path) as img:
        prepared = _preprocess_for_ocr(img.convert("RGB"), lang=lang)
        # two OCR passes; keep the better text
        t1 = pytesseract.image_to_string(prepared, lang=lang, config="--oem 1 --psm 6").strip()
        t2 = pytesseract.image_to_string(prepared, lang=lang, config="--oem 1 --psm 11").strip()
        return t1 if _score_ocr_text(t1) >= _score_ocr_text(t2) else t2


def _ocr_image_easyocr(image_path: Path, lang: str) -> str:
    if easyocr is None:
        raise RuntimeError("easyocr not installed")
    langs = ["en"] if lang == "eng" else [x for x in lang.replace("+", ",").split(",") if x]
    reader = easyocr.Reader(langs or ["en"], gpu=False)
    if Image is not None:
        with Image.open(image_path) as img:
            prepared = _preprocess_for_ocr(img.convert("RGB"), lang=lang)
            tmp = Path(tempfile.gettempdir()) / f"easyocr_pre_{_stamp()}.png"
            prepared.save(tmp)
        try:
            rows = reader.readtext(str(tmp), detail=0, paragraph=True)
        finally:
            tmp.unlink(missing_ok=True)
    else:
        rows = reader.readtext(str(image_path), detail=0, paragraph=True)
    return "\n".join(str(r) for r in rows).strip()


def _extract_image_text(image_path: Path, engine: str, lang: str) -> list[ExtractedPage]:
    if engine == "tesseract":
        text = _ocr_image_tesseract(image_path, lang)
    elif engine == "easyocr":
        text = _ocr_image_easyocr(image_path, lang)
    else:
        if easyocr is not None:
            try:
                text = _ocr_image_easyocr(image_path, lang)
            except Exception:
                text = _ocr_image_tesseract(image_path, lang)
        else:
            text = _ocr_image_tesseract(image_path, lang)
    return [ExtractedPage(source_file=str(image_path), page=1, text=text)]


def _extract_pdf_text_raster_ocr(pdf_path: Path, engine: str, lang: str, *, dpi: int) -> list[ExtractedPage]:
    if fitz is None:
        raise RuntimeError("Missing dependency: pymupdf")
    doc = fitz.open(pdf_path)
    pages: list[ExtractedPage] = []
    try:
        for idx in range(len(doc)):
            page = doc.load_page(idx)
            pix = page.get_pixmap(dpi=max(150, int(dpi)))
            temp_png = Path(tempfile.gettempdir()) / f"{pdf_path.stem}_page_{idx + 1}_{_stamp()}.png"
            pix.save(str(temp_png))
            try:
                extracted = _extract_image_text(temp_png, engine, lang)
                text = extracted[0].text if extracted else ""
            finally:
                temp_png.unlink(missing_ok=True)
            pages.append(ExtractedPage(source_file=str(pdf_path), page=idx + 1, text=text))
    finally:
        doc.close()
    return pages


def _save_text_txt(input_file: Path, pages: list[ExtractedPage], output_dir: Path) -> Path:
    out = _timestamped_path(input_file, "text", "txt", output_dir)
    lines: list[str] = []
    for page in pages:
        lines.append(f"--- Page {page.page} ---")
        lines.append(page.text)
        lines.append("")
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out


def _save_text_json(input_file: Path, pages: list[ExtractedPage], output_dir: Path) -> Path:
    out = _timestamped_path(input_file, "text", "json", output_dir)
    payload = {
        "source_file": str(input_file),
        "pages": [{"page": p.page, "text": p.text} for p in pages],
        "text": "\n\n".join(p.text for p in pages if p.text),
    }
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return out


def _save_text_docx(input_file: Path, pages: list[ExtractedPage], output_dir: Path) -> Path:
    if Document is None:
        raise RuntimeError("Missing dependency: python-docx")
    out = _timestamped_path(input_file, "text", "docx", output_dir)
    doc = Document()
    doc.add_heading(f"Extracted Text: {input_file.name}", level=1)
    for page in pages:
        doc.add_heading(f"Page {page.page}", level=2)
        doc.add_paragraph(page.text or "")
    doc.save(out)
    return out


def _flatten_tables(tables: list[ExtractedTable]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for table in tables:
        if table.dataframe.empty:
            continue
        for _, row in table.dataframe.iterrows():
            rec: dict[str, Any] = {
                "source_file": table.source_file,
                "page": table.page,
                "table_index": table.table_index,
            }
            for key, value in row.to_dict().items():
                rec[str(key)] = value
            rows.append(rec)
    return pd.DataFrame(rows)


def _save_tables_csv(input_file: Path, tables: list[ExtractedTable], output_dir: Path) -> Optional[Path]:
    df = _flatten_tables(tables)
    if df.empty:
        return None
    out = _timestamped_path(input_file, "table", "csv", output_dir)
    df.to_csv(out, index=False)
    return out


def _save_tables_xlsx(input_file: Path, tables: list[ExtractedTable], output_dir: Path) -> Optional[Path]:
    if not tables:
        return None
    out = _timestamped_path(input_file, "table", "xlsx", output_dir)
    with pd.ExcelWriter(out, engine="xlsxwriter") as writer:
        for table in tables:
            sheet = f"P{table.page}_T{table.table_index}"[:31]
            table.dataframe.to_excel(writer, sheet_name=sheet, index=False)
    return out


def _save_tables_json(input_file: Path, tables: list[ExtractedTable], output_dir: Path) -> Optional[Path]:
    if not tables:
        return None
    out = _timestamped_path(input_file, "table", "json", output_dir)
    payload = {
        "source_file": str(input_file),
        "tables": [
            {"page": t.page, "table_index": t.table_index, "rows": t.dataframe.replace(np.nan, None).to_dict(orient="records")}
            for t in tables
        ],
    }
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return out


def _save_tables_docx(input_file: Path, tables: list[ExtractedTable], output_dir: Path) -> Optional[Path]:
    if Document is None:
        raise RuntimeError("Missing dependency: python-docx")
    if not tables:
        return None
    out = _timestamped_path(input_file, "table", "docx", output_dir)
    doc = Document()
    doc.add_heading(f"Extracted Tables: {input_file.name}", level=1)
    for table in tables:
        frame = table.dataframe.fillna("")
        doc.add_heading(f"Page {table.page} - Table {table.table_index}", level=2)
        doc_table = doc.add_table(rows=len(frame) + 1, cols=len(frame.columns))
        for c_idx, col in enumerate(frame.columns):
            doc_table.rows[0].cells[c_idx].text = str(col)
        for r_idx, (_, row) in enumerate(frame.iterrows(), start=1):
            for c_idx, value in enumerate(row.tolist()):
                doc_table.rows[r_idx].cells[c_idx].text = str(value)
        doc.add_paragraph("")
    doc.save(out)
    return out


def run_extract(args: SimpleNamespace) -> SharedResult:
    t0 = perf_counter()

    raw_inputs = list(getattr(args, "input_file", []) or [])
    flattened_inputs = parse_comma_separated_cli_tokens(raw_inputs)
    if not flattened_inputs:
        return failure(
            "Pass at least one -i/--input-file value.",
            exit_code=EXIT_VALIDATION_ERROR,
        )

    mode = str(getattr(args, "mode", "text") or "text").strip().lower()
    out_format = str(getattr(args, "output_format", "json") or "json").strip().lower()
    ocr_mode = str(getattr(args, "ocr", "auto") or "auto").strip().lower()
    image_ocr_engine = str(getattr(args, "image_ocr_engine", "auto") or "auto").strip().lower()
    table_engine = str(getattr(args, "table_engine", "auto") or "auto").strip().lower()
    ocr_pdf_quality = str(getattr(args, "ocr_pdf_quality", "premium") or "premium").strip().lower()
    ocr_pdf_output_type = str(getattr(args, "ocr_pdf_output_type", "pdfa") or "pdfa").strip().lower()
    ocr_threshold = int(getattr(args, "ocr_threshold", 80) or 80)
    image_pdf_dpi = int(getattr(args, "image_pdf_dpi", 450) or 450)
    lang = str(getattr(args, "lang", "eng") or "eng").strip()
    max_files = int(getattr(args, "max_files", 5000) or 5000)
    allow_empty_match = bool(getattr(args, "allow_empty_match", False))
    out_dir_arg = getattr(args, "output_directory", None)
    output_directory = Path(out_dir_arg).expanduser().resolve() if out_dir_arg is not None else None

    if mode not in {"text", "table", "ocrpdf"}:
        return failure("mode must be one of: text, table, ocrpdf.", exit_code=EXIT_VALIDATION_ERROR)
    if ocr_mode not in {"auto", "always", "never"}:
        return failure("ocr must be one of: auto, always, never.", exit_code=EXIT_VALIDATION_ERROR)
    if image_ocr_engine not in {"auto", "tesseract", "easyocr"}:
        return failure("image_ocr_engine must be one of: auto, tesseract, easyocr.", exit_code=EXIT_VALIDATION_ERROR)
    if table_engine not in {"auto", "camelot", "pdfplumber"}:
        return failure("table_engine must be one of: auto, camelot, pdfplumber.", exit_code=EXIT_VALIDATION_ERROR)
    if mode == "text" and out_format not in TEXT_FORMATS:
        return failure(f"mode=text supports formats: {sorted(TEXT_FORMATS)}.", exit_code=EXIT_VALIDATION_ERROR)
    if mode == "table" and out_format not in TABLE_FORMATS:
        return failure(f"mode=table supports formats: {sorted(TABLE_FORMATS)}.", exit_code=EXIT_VALIDATION_ERROR)
    if mode == "ocrpdf" and out_format not in {"pdf", ""}:
        return failure("mode=ocrpdf supports only output format pdf.", exit_code=EXIT_VALIDATION_ERROR)
    if image_pdf_dpi < 150 or image_pdf_dpi > 1200:
        return failure("image_pdf_dpi must be between 150 and 1200.", exit_code=EXIT_VALIDATION_ERROR)

    try:
        input_paths = _resolve_input_paths(flattened_inputs, allow_empty_match=allow_empty_match, max_files=max_files)
    except Exception as exc:  # noqa: BLE001
        return failure(str(exc), exit_code=EXIT_VALIDATION_ERROR, error=str(exc))

    if not input_paths:
        if allow_empty_match:
            return success("No files matched; nothing to do.", outputs=[])
        return failure("No valid input files found.", exit_code=EXIT_VALIDATION_ERROR)

    if mode == "ocrpdf" and shutil.which("ocrmypdf") is None:
        return failure("Missing required system dependency: ocrmypdf", exit_code=EXIT_VALIDATION_ERROR)
    if fitz is None:
        return failure("Missing dependency: pymupdf", exit_code=EXIT_VALIDATION_ERROR)
    if mode == "table" and pdfplumber is None and camelot is None:
        return failure("Missing table extractors: install camelot-py or pdfplumber.", exit_code=EXIT_VALIDATION_ERROR)
    if out_format == "docx" and Document is None:
        return failure("Missing dependency: python-docx", exit_code=EXIT_VALIDATION_ERROR)

    outputs: list[str] = []
    for in_path in input_paths:
        target_dir = output_directory if output_directory is not None else in_path.parent.resolve()
        target_dir.mkdir(parents=True, exist_ok=True)
        temp_ocr_pdf: Optional[Path] = None
        temp_input_pdf: Optional[Path] = None

        try:
            if mode == "ocrpdf":
                source_pdf = in_path
                if in_path.suffix.lower() != ".pdf":
                    source_pdf = _image_to_temp_pdf(in_path, dpi=image_pdf_dpi)
                    temp_input_pdf = source_pdf
                out = _timestamped_path(in_path, "ocrpdf", "pdf", target_dir)
                ok, err = _ocr_pdf(
                    source_pdf,
                    out,
                    lang=lang,
                    quality=ocr_pdf_quality,
                    output_type=ocr_pdf_output_type,
                )
                if not ok:
                    msg = f"OCR PDF generation failed for: {in_path.name}"
                    if err:
                        msg += f" ({err})"
                    return failure(msg, exit_code=EXIT_ERROR)
            elif in_path.suffix.lower() == ".pdf":
                extract_pdf = _maybe_ocr_pdf(
                    in_path,
                    ocr_mode,
                    ocr_threshold,
                    lang=lang,
                    quality=ocr_pdf_quality,
                    output_type=ocr_pdf_output_type,
                )
                if extract_pdf != in_path:
                    temp_ocr_pdf = extract_pdf

                if mode == "text":
                    pages = _extract_pdf_text_layout(extract_pdf)
                    low_text = _text_chars_count(pages) < max(1, ocr_threshold)
                    if ocr_mode != "never" and low_text:
                        pages = _extract_pdf_text_raster_ocr(extract_pdf, image_ocr_engine, lang, dpi=image_pdf_dpi)
                    if out_format == "txt":
                        out = _save_text_txt(in_path, pages, target_dir)
                    elif out_format == "json":
                        out = _save_text_json(in_path, pages, target_dir)
                    else:
                        out = _save_text_docx(in_path, pages, target_dir)
                else:
                    tables = _extract_pdf_tables(extract_pdf, table_engine)
                    if out_format == "csv":
                        out = _save_tables_csv(in_path, tables, target_dir)
                    elif out_format == "xlsx":
                        out = _save_tables_xlsx(in_path, tables, target_dir)
                    elif out_format == "json":
                        out = _save_tables_json(in_path, tables, target_dir)
                    else:
                        out = _save_tables_docx(in_path, tables, target_dir)
            else:
                if mode != "text":
                    continue
                pages = _extract_image_text(in_path, image_ocr_engine, lang)
                if out_format == "txt":
                    out = _save_text_txt(in_path, pages, target_dir)
                elif out_format == "json":
                    out = _save_text_json(in_path, pages, target_dir)
                else:
                    out = _save_text_docx(in_path, pages, target_dir)

            if out is not None:
                outputs.append(str(out))
        except Exception as exc:  # noqa: BLE001
            return failure(f"Extraction failed for {in_path.name}: {exc}", exit_code=EXIT_ERROR, error=str(exc))
        finally:
            if temp_ocr_pdf is not None and temp_ocr_pdf.exists():
                temp_ocr_pdf.unlink(missing_ok=True)
            if temp_input_pdf is not None and temp_input_pdf.exists():
                temp_input_pdf.unlink(missing_ok=True)

    if not outputs:
        return failure("No output generated.", exit_code=EXIT_ERROR)

    duration = perf_counter() - t0
    msg = f"Wrote {len(outputs)} output file(s) in {format_duration(duration)}"
    artifacts = [{"role": "output", "path": p} for p in outputs]
    return success(
        msg,
        mode=mode,
        output_format=out_format,
        outputs=outputs,
        artifacts=artifacts,
        duration_sec=round(duration, 4),
        duration_human=format_duration(duration),
    )
